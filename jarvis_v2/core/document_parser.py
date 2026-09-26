#!/usr/bin/env python3
"""
JARVIS Document Parser - Multi-format document parsing
Supports PDF, DOCX, TXT, MD, HTML
"""

import os
from pathlib import Path
from typing import Dict, List, Optional
import PyPDF2
import docx
import markdown
from bs4 import BeautifulSoup

class DocumentParser:
    """
    Multi-format document parser
    Extracts text from various document formats
    """
    
    SUPPORTED_FORMATS = ['.pdf', '.docx', '.txt', '.md', '.html', '.htm']
    
    def __init__(self):
        """Initialize document parser"""
        print("📄 Document Parser initialized")
    
    def parse(self, file_path: str) -> Dict:
        """
        Parse document and extract text
        
        Args:
            file_path: Path to document
        
        Returns:
            Dict with:
                - text: Extracted text
                - metadata: Document metadata
                - pages: List of page texts (if applicable)
        """
        path = Path(file_path)
        
        if not path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")
        
        extension = path.suffix.lower()
        
        if extension not in self.SUPPORTED_FORMATS:
            raise ValueError(f"Unsupported format: {extension}")
        
        # Route to appropriate parser
        if extension == '.pdf':
            return self._parse_pdf(path)
        elif extension == '.docx':
            return self._parse_docx(path)
        elif extension == '.txt':
            return self._parse_txt(path)
        elif extension == '.md':
            return self._parse_markdown(path)
        elif extension in ['.html', '.htm']:
            return self._parse_html(path)
        else:
            raise ValueError(f"No parser for: {extension}")
    
    def _parse_pdf(self, path: Path) -> Dict:
        """Parse PDF document"""
        try:
            with open(path, 'rb') as file:
                reader = PyPDF2.PdfReader(file)
                
                # Extract metadata
                metadata = {
                    'source': path.name,
                    'format': 'pdf',
                    'pages': len(reader.pages),
                    'title': reader.metadata.title if reader.metadata else None,
                    'author': reader.metadata.author if reader.metadata else None
                }
                
                # Extract text from each page
                pages = []
                full_text = []
                
                for i, page in enumerate(reader.pages):
                    page_text = page.extract_text()
                    pages.append({
                        'page_num': i + 1,
                        'text': page_text
                    })
                    full_text.append(page_text)
                
                return {
                    'text': '\n\n'.join(full_text),
                    'metadata': metadata,
                    'pages': pages
                }
                
        except Exception as e:
            raise Exception(f"Error parsing PDF: {e}")
    
    def _parse_docx(self, path: Path) -> Dict:
        """Parse DOCX document"""
        try:
            doc = docx.Document(path)
            
            # Extract metadata
            metadata = {
                'source': path.name,
                'format': 'docx',
                'paragraphs': len(doc.paragraphs)
            }
            
            # Extract text from paragraphs
            paragraphs = []
            full_text = []
            
            for i, para in enumerate(doc.paragraphs):
                if para.text.strip():
                    paragraphs.append({
                        'para_num': i + 1,
                        'text': para.text
                    })
                    full_text.append(para.text)
            
            return {
                'text': '\n\n'.join(full_text),
                'metadata': metadata,
                'pages': paragraphs  # Use paragraphs as "pages"
            }
            
        except Exception as e:
            raise Exception(f"Error parsing DOCX: {e}")
    
    def _parse_txt(self, path: Path) -> Dict:
        """Parse plain text document"""
        try:
            with open(path, 'r', encoding='utf-8') as file:
                text = file.read()
            
            # Split into paragraphs
            paragraphs = [p.strip() for p in text.split('\n\n') if p.strip()]
            
            metadata = {
                'source': path.name,
                'format': 'txt',
                'paragraphs': len(paragraphs)
            }
            
            pages = [{'para_num': i + 1, 'text': p} for i, p in enumerate(paragraphs)]
            
            return {
                'text': text,
                'metadata': metadata,
                'pages': pages
            }
            
        except Exception as e:
            raise Exception(f"Error parsing TXT: {e}")
    
    def _parse_markdown(self, path: Path) -> Dict:
        """Parse Markdown document"""
        try:
            with open(path, 'r', encoding='utf-8') as file:
                md_text = file.read()
            
            # Convert to HTML then extract text
            html = markdown.markdown(md_text)
            soup = BeautifulSoup(html, 'html.parser')
            text = soup.get_text()
            
            # Split into sections (by headers)
            sections = []
            current_section = []
            
            for line in md_text.split('\n'):
                if line.startswith('#'):
                    if current_section:
                        sections.append('\n'.join(current_section))
                    current_section = [line]
                else:
                    current_section.append(line)
            
            if current_section:
                sections.append('\n'.join(current_section))
            
            metadata = {
                'source': path.name,
                'format': 'markdown',
                'sections': len(sections)
            }
            
            pages = [{'section_num': i + 1, 'text': s} for i, s in enumerate(sections)]
            
            return {
                'text': text,
                'metadata': metadata,
                'pages': pages
            }
            
        except Exception as e:
            raise Exception(f"Error parsing Markdown: {e}")
    
    def _parse_html(self, path: Path) -> Dict:
        """Parse HTML document"""
        try:
            with open(path, 'r', encoding='utf-8') as file:
                html = file.read()
            
            soup = BeautifulSoup(html, 'html.parser')
            
            # Remove script and style elements
            for script in soup(["script", "style"]):
                script.decompose()
            
            # Extract text
            text = soup.get_text()
            
            # Clean up whitespace
            lines = (line.strip() for line in text.splitlines())
            chunks = (phrase.strip() for line in lines for phrase in line.split("  "))
            text = '\n'.join(chunk for chunk in chunks if chunk)
            
            metadata = {
                'source': path.name,
                'format': 'html',
                'title': soup.title.string if soup.title else None
            }
            
            # Split into paragraphs
            paragraphs = [p.strip() for p in text.split('\n\n') if p.strip()]
            pages = [{'para_num': i + 1, 'text': p} for i, p in enumerate(paragraphs)]
            
            return {
                'text': text,
                'metadata': metadata,
                'pages': pages
            }
            
        except Exception as e:
            raise Exception(f"Error parsing HTML: {e}")
    
    def is_supported(self, file_path: str) -> bool:
        """Check if file format is supported"""
        extension = Path(file_path).suffix.lower()
        return extension in self.SUPPORTED_FORMATS

def main():
    """Test document parser"""
    print("🧪 Testing Document Parser...")
    print("=" * 60)
    
    parser = DocumentParser()
    
    # Test with sample text
    print("\n1. Creating test document...")
    test_file = Path("test_document.txt")
    test_content = """JARVIS AI Assistant

This is a test document for the document parser.

It contains multiple paragraphs to test the parsing functionality.

The parser should extract all text and metadata correctly."""
    
    with open(test_file, 'w') as f:
        f.write(test_content)
    
    print(f"   ✓ Created {test_file}")
    
    # Parse document
    print("\n2. Parsing document...")
    result = parser.parse(str(test_file))
    
    print(f"   Format: {result['metadata']['format']}")
    print(f"   Paragraphs: {result['metadata']['paragraphs']}")
    print(f"   Text length: {len(result['text'])} chars")
    print(f"   Pages/sections: {len(result['pages'])}")
    
    # Clean up
    test_file.unlink()
    print(f"\n   ✓ Cleaned up test file")
    
    print("\n" + "=" * 60)
    print("✓ Document Parser test complete")

if __name__ == "__main__":
    main()
